"""BigStep branded DOCX — comprehensive sample with every component including
normal-weight bullet lists. EMBEDDED Poppins."""
import sys, os
sys.path.insert(0, "../scripts")
import bigstep_brand as bs
from bigstep_brand import C, BRAND
from embed_fonts import embed_fonts_in_docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Poppins"
def rgb(h): return RGBColor(*bs.hx(h))

def _rfonts(rpr, name):
    rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a), name)

def set_font(run, size, color, bold=False):
    run.font.name=FONT; run.font.size=Pt(size); run.font.color.rgb=rgb(color); run.font.bold=bold
    _rfonts(run._element.get_or_add_rPr(), FONT)

def set_doc_default_font(doc, name):
    el=doc.styles.element; dd=el.find(qn('w:docDefaults'))
    if dd is None: dd=OxmlElement('w:docDefaults'); el.insert(0,dd)
    rprd=dd.find(qn('w:rPrDefault'))
    if rprd is None: rprd=OxmlElement('w:rPrDefault'); dd.append(rprd)
    rpr=rprd.find(qn('w:rPr'))
    if rpr is None: rpr=OxmlElement('w:rPr'); rprd.append(rpr)
    _rfonts(rpr,name)
    n=doc.styles['Normal']; n.font.name=name; _rfonts(n.element.get_or_add_rPr(),name)

def shade(cell,h):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),h.lstrip('#')); tcPr.append(sh)

def left_accent(cell,sz,color):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        e=OxmlElement(f'w:{edge}')
        if edge=='left': e.set(qn('w:val'),'single'); e.set(qn('w:sz'),str(sz)); e.set(qn('w:color'),color.lstrip('#'))
        else: e.set(qn('w:val'),'nil')
        b.append(e)
    tcPr.append(b)

def cell_pad(cell,dxa=140):
    tcPr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for m in ('top','bottom','start','end'):
        e=OxmlElement(f'w:{m}'); e.set(qn('w:w'),str(dxa)); e.set(qn('w:type'),'dxa'); mar.append(e)
    tcPr.append(mar)

def no_borders(t):
    tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'nil'); b.append(e)
    tblPr.append(b)

def para(doc,runs,after=6,before=0,align=WD_ALIGN_PARAGRAPH.LEFT,line=1.2):
    p=doc.add_paragraph(); p.alignment=align
    pf=p.paragraph_format; pf.space_after=Pt(after); pf.space_before=Pt(before); pf.line_spacing=line
    for (t,s,c,b) in runs: set_font(p.add_run(t),s,c,b)
    return p

def bullets(doc, items, after=4, size=10.5):
    """Normal-weight body-text bullet list. Items are NEVER bold — bullets use
    the same Poppins Regular / slate / 10.5pt as body paragraphs."""
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.2
        # clear any inherited bold from the style and set brand body font
        for r in p.runs:
            set_font(r, size, C.SLATE, False)
        if not p.runs:
            set_font(p.add_run(item), size, C.SLATE, False)
        else:
            p.runs[0].text = item
            set_font(p.runs[0], size, C.SLATE, False)

def two_tone_title(doc,lead,tail,size=24,before=4):
    return para(doc,[(lead,size,C.NAVY,True),(tail,size,C.PRIMARY,True)],after=2,before=before)

def accent_rule(doc,width_in=1.5):
    t=doc.add_table(rows=1,cols=1); t.autofit=False
    t.columns[0].width=Inches(width_in); t.rows[0].height=Pt(2.5)
    no_borders(t); shade(t.rows[0].cells[0],C.PRIMARY)
    r=t.rows[0].cells[0].paragraphs[0]; r.paragraph_format.space_after=Pt(0); set_font(r.add_run(" "),2,C.PRIMARY)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def section_head(doc,lead,tail,before=6):
    two_tone_title(doc,lead,tail,size=18,before=before); accent_rule(doc,1.2)

def header_bottom_border(p,color,sz):
    pPr=p._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),str(sz)); bottom.set(qn('w:space'),'4'); bottom.set(qn('w:color'),color.lstrip('#'))
    pbdr.append(bottom); pPr.append(pbdr)

def stat_row(doc,stats):
    st=doc.add_table(rows=1,cols=len(stats)); st.alignment=WD_TABLE_ALIGNMENT.CENTER; st.autofit=False
    no_borders(st)
    for j,(v,l,a) in enumerate(stats):
        cell=st.rows[0].cells[j]; cell.width=Inches(6.6/len(stats))
        shade(cell,C.TINT); left_accent(cell,24,a); cell_pad(cell)
        p1=cell.paragraphs[0]; p1.paragraph_format.space_after=Pt(2); set_font(p1.add_run(v),20,C.NAVY,True)
        set_font(cell.add_paragraph().add_run(l),9,C.SLATE,False)
    doc.add_paragraph().paragraph_format.space_after=Pt(6)

def callout(doc,title,body,accent=None,fill=None):
    fill=fill or C.TINT; accent=accent or C.PRIMARY
    t=doc.add_table(rows=1,cols=1); t.autofit=False; t.columns[0].width=Inches(6.6)
    no_borders(t); cell=t.rows[0].cells[0]; shade(cell,fill); left_accent(cell,30,accent); cell_pad(cell,180)
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3); set_font(p.add_run(title),11,C.NAVY,True)
    for ln in body:
        pp=cell.add_paragraph(); pp.paragraph_format.space_after=Pt(2); set_font(pp.add_run(ln),9.5,C.SLATE,False)
    doc.add_paragraph().paragraph_format.space_after=Pt(4)

def two_col_cards(doc,items):
    rows=(len(items)+1)//2
    t=doc.add_table(rows=rows,cols=2); t.autofit=False; no_borders(t)
    for idx,(title,body) in enumerate(items):
        cell=t.rows[idx//2].cells[idx%2]; cell.width=Inches(3.25)
        shade(cell,C.TINT); left_accent(cell,24,C.PRIMARY if idx%2==0 else C.CYAN); cell_pad(cell,150)
        p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); set_font(p.add_run(title),11,C.NAVY,True)
        set_font(cell.add_paragraph().add_run(body),9.5,C.SLATE,False)
    doc.add_paragraph().paragraph_format.space_after=Pt(6)

def branded_table(doc,headers,rows,widths):
    bt=doc.add_table(rows=1,cols=len(headers)); bt.autofit=False
    for j,h in enumerate(headers):
        c=bt.rows[0].cells[j]; c.width=widths[j]; shade(c,C.PRIMARY); cell_pad(c,120)
        set_font(c.paragraphs[0].add_run(h),10,C.WHITE,True)
    for i,row in enumerate(rows):
        cells=bt.add_row().cells; fillc=C.WHITE if i%2==0 else C.TINT2
        for j,val in enumerate(row):
            cells[j].width=widths[j]; shade(cells[j],fillc); cell_pad(cells[j],120)
            set_font(cells[j].paragraphs[0].add_run(val),10,C.SLATE,False)
    tblPr=bt._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),C.TINT.lstrip('#')); b.append(e)
    tblPr.append(b)
    doc.add_paragraph().paragraph_format.space_after=Pt(6)

def contact_banner(doc):
    t=doc.add_table(rows=1,cols=1); t.autofit=False; t.columns[0].width=Inches(6.6)
    no_borders(t); cell=t.rows[0].cells[0]; shade(cell,C.PRIMARY); cell_pad(cell,220)
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(4)
    r=p.add_run(); r.add_picture(bs.logo_path("white"), width=Inches(1.6))
    set_font(cell.add_paragraph().add_run("Let's build your next big step."),16,C.WHITE,True)
    cp=cell.add_paragraph(); cp.paragraph_format.space_before=Pt(4)
    set_font(cp.add_run(BRAND["email"]+"   \u00b7   "+BRAND["website"]),11,C.WHITE,False)
    set_font(cell.add_paragraph().add_run("Founded "+str(BRAND["founded"])+"  \u00b7  "+BRAND["hq"]),9,"#D9E6FF",False)

def page_break(doc): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ----------------------------------------------------------------- build
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.95); sec.bottom_margin=Inches(0.7); sec.left_margin=Inches(0.9); sec.right_margin=Inches(0.9)
set_doc_default_font(doc,FONT)
doc.styles['Normal'].font.size=Pt(10.5); doc.styles['Normal'].font.color.rgb=rgb(C.SLATE)
# Force List Bullet style to normal weight
if 'List Bullet' in doc.styles:
    lb=doc.styles['List Bullet']; lb.font.bold=False; lb.font.name=FONT; lb.font.size=Pt(10.5); lb.font.color.rgb=rgb(C.SLATE)
    _rfonts(lb.element.get_or_add_rPr(), FONT)

hp=sec.header.paragraphs[0]; hp.add_run().add_picture(bs.logo_path('color'),width=Inches(0.7))
# blue accent only — no grey line (removed header_bottom_border)
fp=sec.footer.paragraphs[0]; set_font(fp.add_run(BRAND['website']),8,C.MUTED)

# PAGE 1 — overview
two_tone_title(doc,"Capability & Engagement ","Overview",size=24)
accent_rule(doc,1.5)
para(doc,[("Company Profile  \u00b7  2026  \u00b7  Confidential",9.5,C.MUTED,False)],after=10)
para(doc,[("BigStep Technologies is an AI-First, Cloud-Native software engineering and product "
           "development partner. We embed senior, cross-functional squads inside your team to ship "
           "resilient products faster \u2014 from first prototype to global scale.",10.5,C.SLATE,False)],after=10)
stat_row(doc,[("17+","Years delivering software & AI",C.PRIMARY),
              ("700+","Apps, products & solutions shipped",C.CYAN),
              ("500+","Customers across 6+ industries",C.PRIMARY)])
section_head(doc,"How We ","Deliver")
callout(doc,"Squad-as-a-Service",
        ["Dedicated, hand-picked developers, designers, QA and DevOps that work as an extension of "
         "your in-house team \u2014 with shared ownership of outcomes and a bias for action.",
         "Scale up or down on demand, matched to your scope, timeline, and risk appetite."],
        accent=C.PRIMARY)

# PAGE 2 — services
page_break(doc)
section_head(doc,"Our Core ","Services",before=2)
para(doc,[("Four practice areas, delivered by one integrated team.",10.5,C.SLATE,False)],after=8)
two_col_cards(doc,[
    ("Digital Products","Product development, web & mobile apps, backend engineering, and QA."),
    ("Artificial Intelligence","Generative AI, agentic AI, and applied ML across your workflows."),
    ("Data Solutions","Data engineering, modernization, and analytics that compound over time."),
    ("Cloud & DevOps","Cloud architecture, managed services, CI/CD, and infra modernization."),
])
section_head(doc,"Key ","Specializations")
bullets(doc, [
    "Agentic workflows and autonomous AI agents",
    "Data extraction and document intelligence",
    "Live video and media streaming at scale",
    "PropTech platforms and construction-tech",
    "SaaS architecture and multi-tenant systems",
    "AWS consulting and cloud migration",
])

# PAGE 3 — outcomes
page_break(doc)
section_head(doc,"Outcomes That ","Compound",before=2)
para(doc,[("Representative results from recent partnerships across PropTech, media streaming, "
           "and healthcare.",10.5,C.SLATE,False)],after=8)
branded_table(doc,["Programme","Focus","Headline Outcome"],
    [["VTS Rise","PropTech platform","Scaled the platform 5X"],
     ["Jio Media Stack","Live streaming SDK","-15% CPU, 500+ streams/channel"],
     ["House of Diagnostics","Healthcare platform","Major lift in online bookings"],
     ["Channelize.io","Live video commerce","70% serverless, faster scaling"],
     ["FitPass","Healthtech voice + AI","-60% manual documentation"]],
    [Inches(2.0),Inches(2.0),Inches(2.6)])
section_head(doc,"What Partners ","Value")
bullets(doc, [
    "Squads led by engineers and designers who have shipped at scale \u2014 not ramp-up resources",
    "Every deliverable carries a consistent, premium BigStep identity",
    "Transparent communication, weekly demos, and shared ownership of outcomes",
    "Deep domain expertise across PropTech, media, healthtech, and fintech",
])

# PAGE 4 — engagement models
page_break(doc)
section_head(doc,"Engagement ","Models",before=2)
para(doc,[("Pick the model that fits your scope, timeline, and risk appetite.",10.5,C.SLATE,False)],after=8)
branded_table(doc,["Model","What it is","Best for"],
    [["Staff Augmentation","Add skilled engineers to your existing team","Faster product delivery"],
     ["Time & Material","Billing based on actual time & effort","Evolving, iterative work"],
     ["Global Capability Centre","A long-term dedicated team","Multi-year scaling"],
     ["Fixed Cost","Defined scope, timeline & budget","Well-defined projects"],
     ["Hybrid Delivery","Coordinated teams across locations","Speed + cost efficiency"]],
    [Inches(1.9),Inches(2.7),Inches(2.0)])
section_head(doc,"The BigStep ","Difference")
bullets(doc, [
    "Integrity \u2014 we do what we say, every time",
    "Innovation \u2014 we push for better, not just faster",
    "Ownership \u2014 your outcomes are our outcomes",
    "Excellence \u2014 senior-by-default, quality-first delivery",
    "Customer centricity \u2014 we listen before we build",
    "Continuous learning \u2014 every sprint makes the team sharper",
    "Bias for action \u2014 surface risks early, ship often",
    "Collaboration \u2014 one team, one backlog, one goal",
])

# PAGE 5 — contact
page_break(doc)
section_head(doc,"Let's ","Talk",before=2)
para(doc,[("Interested in a detailed profile for any of our services? We'd be glad to share.",10.5,C.SLATE,False)],after=10)
contact_banner(doc)

out=os.environ.get("OUT","/mnt/user-data/outputs/BigStep_Brand_Sample.docx")
doc.save(out)
embed_fonts_in_docx(out,[{"name":"Poppins",
    "regular":str(bs.FONTS/"Poppins-Regular.ttf"),
    "bold":str(bs.FONTS/"Poppins-Bold.ttf")}])
print("saved + fonts embedded:",out)
