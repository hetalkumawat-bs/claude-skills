"""BigStep branded Excel workbook — production-grade, blue-only, Poppins.
Every sheet carries a consistent branded header (logo + two-tone title + cyan
accent + section label) and a printed footer (website + page). Branded tables,
live formulas (zero errors), and a brand-blue chart."""
import sys, os
sys.path.insert(0, "../scripts")
import bigstep_brand as bs
from bigstep_brand import C, BRAND
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

def hexa(h): return h.lstrip('#').upper()
NAVY=hexa(C.NAVY); PRIMARY=hexa(C.PRIMARY); CYAN=hexa(C.CYAN); SLATE=hexa(C.SLATE)
TINT=hexa(C.TINT); TINT2=hexa(C.TINT2); WHITE="FFFFFF"; MUTED=hexa(C.MUTED)
FONT="Poppins"
ASSETS=bs.ASSETS
thin=Side(style="thin", color=TINT)
GRID=Border(left=thin,right=thin,top=thin,bottom=thin)

def cell(ws, ref, val, size=10.5, color=SLATE, bold=False, fill=None, align="left",
         wrap=False, border=False, numfmt=None, valign="center"):
    c=ws[ref]; c.value=val
    c.font=Font(name=FONT, size=size, color=color, bold=bold)
    c.alignment=Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    if fill: c.fill=PatternFill(fill_type="solid", fgColor=fill)
    if border: c.border=GRID
    if numfmt: c.number_format=numfmt
    return c

def img(path, h_px):
    im=XLImage(path); ratio=im.width/im.height; im.height=h_px; im.width=int(h_px*ratio); return im

def sheet_header(ws, lead, tail, label, last_col="F"):
    ws.sheet_view.showGridLines=False
    ws.column_dimensions["A"].width=2.4
    ws.row_dimensions[1].height=6
    ws.row_dimensions[2].height=30
    ws.row_dimensions[3].height=4
    ws.row_dimensions[4].height=24
    ws.row_dimensions[5].height=8
    ws.row_dimensions[6].height=8
    # logo
    ws.add_image(img(bs.logo_path("color"), 34), "B2")
    # two-tone title (rich text)
    ws["B4"]=CellRichText([
        TextBlock(InlineFont(rFont=FONT, sz=18, b=True, color=NAVY), lead),
        TextBlock(InlineFont(rFont=FONT, sz=18, b=True, color=PRIMARY), tail)])
    ws["B4"].alignment=Alignment(horizontal="left", vertical="center")
    # section label, top-right beside the logo (avoids colliding with the title)
    cc=ws.cell(row=2, column=__import__("openpyxl").utils.column_index_from_string(last_col))
    cc.value=label.upper(); cc.font=Font(name=FONT, size=9, color=MUTED, bold=True)
    cc.alignment=Alignment(horizontal="right", vertical="center")
    # short cyan accent bar under title (anchored image)
    acc=img(str(ASSETS/"accent_cyan.png"), 5); acc.width=42
    ws.add_image(acc, "B5")
    # faint full-width rule (bottom border on row 6 across content cols)
    from openpyxl.utils import column_index_from_string
    for ci in range(2, column_index_from_string(last_col)+1):
        cc=ws.cell(row=6, column=ci)
        cc.border=Border(bottom=Side(style="thin", color=TINT))

def set_footer(ws):
    ws.oddFooter.left.text  = '&"Poppins"&8&K6B7A8D www.bigsteptech.com'
    ws.oddFooter.right.text = '&"Poppins"&8&K6B7A8D Page &P of &N '
    ws.oddFooter.center.text = ""

def page_setup(ws, area):
    ws.page_setup.orientation="portrait"
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
    ws.page_margins.left=ws.page_margins.right=0.4
    ws.page_margins.top=0.5; ws.page_margins.bottom=0.6
    ws.print_area=area

def branded_table(ws, top, headers, rows, widths, start_col=2):
    sc=start_col
    for j,h in enumerate(headers):
        cell(ws, f"{get_column_letter(sc+j)}{top}", h, size=11, color=WHITE, bold=True,
             fill=PRIMARY, align="left", border=True)
    ws.row_dimensions[top].height=26
    for i,row in enumerate(rows):
        rr=top+1+i; fillc=WHITE if i%2==0 else TINT2
        for j,val in enumerate(row):
            cell(ws, f"{get_column_letter(sc+j)}{rr}", val, size=10.5, color=SLATE,
                 fill=fillc, align="left", border=True, wrap=True)
        ws.row_dimensions[rr].height=24
    for j,w in enumerate(widths):
        ws.column_dimensions[get_column_letter(sc+j)].width=w
    return top+1+len(rows)

wb=openpyxl.Workbook()

# ===================== Sheet 1: Overview =====================
ws=wb.active; ws.title="Overview"
sheet_header(ws,"Capability & Engagement ","Overview","Company Profile · 2026", last_col="E")
cell(ws,"B8","BigStep Technologies is an AI-First, Cloud-Native software engineering and product "
     "development partner \u2014 embedding senior, cross-functional squads inside your team to ship "
     "resilient products faster, from first prototype to global scale.",
     size=10.5, color=SLATE, wrap=True, valign="top")
ws.merge_cells("B8:E11")
for rr in (8,9,10,11): ws.row_dimensions[rr].height=15
cell(ws,"B13","At a Glance", size=12, color=NAVY, bold=True)
endr=branded_table(ws, 14, ["Metric","Value"],
    [["Years of experience","17+"],
     ["Apps, products & solutions delivered","700+"],
     ["Customers worldwide","500+"],
     ["BigStep'ians","180+"],
     ["Industries served","6+"],
     ["Positive online reviews","100+"]],
    [44,16])
cell(ws,"D12","",); 
page_setup(ws, f"A1:E{endr+1}"); set_footer(ws)

# ===================== Sheet 2: Delivery Metrics =====================
ws2=wb.create_sheet("Delivery Metrics")
sheet_header(ws2,"Delivery ","Metrics","Illustrative sample data", last_col="E")
top=8
endr=branded_table(ws2, top, ["Year","Apps Delivered","Customers Added"],
    [["2022",120,60],["2023",165,82],["2024",210,104],["2025",205,118]],
    [16,20,20])
tr=endr
cell(ws2,f"B{tr}","Total", size=10.5, color=NAVY, bold=True, fill=TINT, border=True)
cell(ws2,f"C{tr}",f"=SUM(C{top+1}:C{tr-1})", size=10.5, color=NAVY, bold=True, fill=TINT, border=True, numfmt="#,##0")
cell(ws2,f"D{tr}",f"=SUM(D{top+1}:D{tr-1})", size=10.5, color=NAVY, bold=True, fill=TINT, border=True, numfmt="#,##0")
ws2.row_dimensions[tr].height=24
# chart (brand-blue bars)
chart=BarChart(); chart.type="col"; chart.title="Apps Delivered by Year"
chart.height=6.6; chart.width=12; chart.legend=None; chart.gapWidth=60
data=Reference(ws2, min_col=3, min_row=top, max_row=tr-1)
cats=Reference(ws2, min_col=2, min_row=top+1, max_row=tr-1)
chart.add_data(data, titles_from_data=True); chart.set_categories(cats)
chart.series[0].graphicalProperties.solidFill=PRIMARY
chart.series[0].graphicalProperties.line.noFill=True
ws2.add_chart(chart, f"B{tr+2}")
page_setup(ws2, f"A1:E{tr+20}"); set_footer(ws2)

# ===================== Sheet 3: Case Studies =====================
ws3=wb.create_sheet("Case Studies")
sheet_header(ws3,"Outcomes That ","Compound","Case Studies", last_col="F")
endr=branded_table(ws3, 8, ["Programme","Focus","Headline Outcome"],
    [["VTS Rise","PropTech platform","Scaled the platform 5X"],
     ["Jio Media Stack","Live streaming SDK","-15% CPU, 500+ streams/channel"],
     ["House of Diagnostics","Healthcare platform","Major lift in online bookings"],
     ["Channelize.io","Live video commerce","70% serverless, faster scaling"],
     ["FitPass","Healthtech voice + AI","-60% manual documentation"]],
    [24,24,38])
page_setup(ws3, f"A1:F{endr+1}"); set_footer(ws3)

# ===================== Sheet 4: Engagement Models =====================
ws4=wb.create_sheet("Engagement Models")
sheet_header(ws4,"Engagement ","Models","How We Engage", last_col="F")
endr=branded_table(ws4, 8, ["Model","What it is","Best for"],
    [["Staff Augmentation","Add skilled engineers to your existing team","Faster product delivery"],
     ["Time & Material","Billing based on actual time & effort","Evolving, iterative work"],
     ["Global Capability Centre","A long-term dedicated team","Multi-year scaling"],
     ["Fixed Cost","Defined scope, timeline & budget","Well-defined projects"],
     ["Hybrid Delivery","Coordinated teams across locations","Speed + cost efficiency"]],
    [26,40,26])
page_setup(ws4, f"A1:F{endr+1}"); set_footer(ws4)

out=os.environ.get("OUT","/mnt/user-data/outputs/BigStep_Brand_Sample.xlsx")
wb.save(out); print("saved", out)
